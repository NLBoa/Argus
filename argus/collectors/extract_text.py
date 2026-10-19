"""Turn course files into text per page or slide, in memory. Nothing is written to disk."""

from __future__ import annotations

import io

from argus.collectors.text import html_to_text
from argus.domain.models import MaterialPage

MAX_BYTES = 30 * 1024 * 1024
PDF_TYPES = ("application/pdf",)
PPTX_TYPES = ("application/vnd.openxmlformats-officedocument.presentationml.presentation",)


def kind_for(content_type: str, filename: str) -> str | None:
    name = filename.lower()
    if content_type in PDF_TYPES or name.endswith(".pdf"):
        return "pdf"
    if content_type in PPTX_TYPES or name.endswith(".pptx"):
        return "pptx"
    return None


def pdf_pages(data: bytes) -> tuple[MaterialPage, ...]:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    out = []
    for n, page in enumerate(reader.pages, start=1):
        text = " ".join((page.extract_text() or "").split())
        if text:
            out.append(MaterialPage(f"page {n}", text))
    return tuple(out)


def pptx_slides(data: bytes) -> tuple[MaterialPage, ...]:
    from pptx import Presentation

    deck = Presentation(io.BytesIO(data))
    out = []
    for n, slide in enumerate(deck.slides, start=1):
        parts: list[str] = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                parts.extend(p.text for p in shape.text_frame.paragraphs if p.text.strip())
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                parts.append(f"(speaker notes) {notes}")
        text = " ".join(" ".join(parts).split())
        if text:
            out.append(MaterialPage(f"slide {n}", text))
    return tuple(out)


def extract(kind: str, data: bytes) -> tuple[MaterialPage, ...]:
    if len(data) > MAX_BYTES:
        raise ValueError("file too large to extract")
    if kind == "pdf":
        return pdf_pages(data)
    if kind == "pptx":
        return pptx_slides(data)
    raise ValueError(f"cannot extract {kind}")


def page_text(html: str | None) -> tuple[MaterialPage, ...]:
    text = html_to_text(html)
    return (MaterialPage("page", text),) if text else ()
